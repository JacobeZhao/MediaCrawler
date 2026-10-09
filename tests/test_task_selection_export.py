import io
import os
import tempfile
import unittest

import aiosqlite
import openpyxl

from service import service_db as db
from service.services import export_service


class SelectedTaskExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.service_path = db.SERVICE_DB_PATH
        self.content_path = export_service.SQLITE_DB_PATH
        db.SERVICE_DB_PATH = os.path.join(self.temp_dir.name, "service.db")
        export_service.SQLITE_DB_PATH = os.path.join(self.temp_dir.name, "content.db")
        await db.init_service_db()
        self.first = await db.create_task(db.TaskType.NOTE, {"notes": []})
        self.second = await db.create_task(db.TaskType.NOTE, {"notes": [1]})
        self.other = await db.create_task(db.TaskType.NOTE, {"notes": [2]})
        async with aiosqlite.connect(export_service.SQLITE_DB_PATH) as connection:
            await connection.execute(
                "CREATE TABLE xhs_note (note_id TEXT PRIMARY KEY, title TEXT, desc TEXT, "
                "nickname TEXT, user_id TEXT, time INTEGER, liked_count TEXT, "
                "comment_count TEXT, collected_count TEXT, note_url TEXT, image_list TEXT)"
            )
            await connection.execute(
                "CREATE TABLE xhs_note_comment (comment_id TEXT PRIMARY KEY, note_id TEXT, "
                "content TEXT, nickname TEXT, user_id TEXT, create_time INTEGER, "
                "like_count TEXT, parent_comment_id TEXT, pictures TEXT)"
            )
            await connection.execute(
                "CREATE TABLE xhs_content_source (entity_type TEXT, entity_id TEXT, "
                "provider TEXT, task_id INTEGER, source_keyword TEXT)"
            )
            await connection.executemany(
                "INSERT INTO xhs_note (note_id, title) VALUES (?,?)",
                [("shared", "=HYPERLINK(\"bad\")"), ("other", "Not selected")],
            )
            await connection.executemany(
                "INSERT INTO xhs_note_comment (comment_id, note_id, content) VALUES (?,?,?)",
                [("comment-1", "shared", "@unsafe\x00"), ("comment-2", "other", "unselected")],
            )
            await connection.executemany(
                "INSERT INTO xhs_content_source VALUES (?,?,?,?,?)",
                [
                    ("note", "shared", "local", self.first, ""),
                    ("note", "shared", "local", self.second, ""),
                    ("comment", "comment-1", "local", self.second, ""),
                    ("note", "other", "local", self.other, ""),
                    ("comment", "comment-2", "local", self.other, ""),
                ],
            )
            await connection.commit()
        self.exporter = export_service.ExportService(os.path.join(self.temp_dir.name, "images"))

    async def asyncTearDown(self):
        db.SERVICE_DB_PATH = self.service_path
        export_service.SQLITE_DB_PATH = self.content_path
        self.temp_dir.cleanup()

    async def test_exports_only_selected_attributions_and_deduplicates(self):
        result = await self.exporter.build_selected_tasks_workbook([self.first, self.second])
        workbook = openpyxl.load_workbook(io.BytesIO(result.read()))
        self.assertEqual(["notes", "comments"], workbook.sheetnames)
        notes = workbook["notes"]
        comments = workbook["comments"]
        self.assertEqual(2, notes.max_row)
        self.assertEqual(2, comments.max_row)
        self.assertEqual("shared", notes["B2"].value)
        self.assertEqual("'=HYPERLINK(\"bad\")", notes["C2"].value)
        self.assertEqual("comment-1", comments["B2"].value)
        self.assertEqual("'@unsafe", comments["D2"].value)
        self.assertEqual({str(self.first), str(self.second)}, set(notes["A2"].value.split(",")))

    async def test_missing_task_or_no_attribution_is_not_exported(self):
        with self.assertRaises(LookupError):
            await self.exporter.build_selected_tasks_workbook([999999])
        no_data = await db.create_task(db.TaskType.NOTE, {"notes": [3]})
        with self.assertRaises(FileNotFoundError):
            await self.exporter.build_selected_tasks_workbook([no_data])


if __name__ == "__main__":
    unittest.main()
