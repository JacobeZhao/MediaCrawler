import json
import os
import site
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]

CHILD_PRELUDE = r'''
import asyncio
import io
import json
import os
import sys
import types
from pathlib import Path
from unittest.mock import AsyncMock, patch

aiosqlite = types.ModuleType("aiosqlite")
aiosqlite.Connection = object
aiosqlite.Row = object()
aiosqlite.IntegrityError = type("IntegrityError", (Exception,), {})
def unexpected_connect(*args, **kwargs):
    raise AssertionError("real database access is forbidden")
aiosqlite.connect = unexpected_connect
sys.modules["aiosqlite"] = aiosqlite

from service.services import export_service

temp_root = Path(sys.argv[1]).resolve()
'''


class ExportServiceContractTests(unittest.TestCase):
    def _run_case(self, body):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            env_file = temp_root / "empty.env"
            env_file.write_bytes(b"")
            pycache = temp_root / "pycache"
            pycache.mkdir()

            import_paths = [str(ROOT)]
            import_paths.extend(path for path in sys.path if path)
            import_paths.extend(site.getsitepackages())
            env = {
                "PYTHONPATH": os.pathsep.join(dict.fromkeys(import_paths)),
                "PYTHONPYCACHEPREFIX": str(pycache),
                "PYTHONDONTWRITEBYTECODE": "1",
                "XHS_ENV_FILE": str(env_file),
            }
            for name in ("SystemRoot", "WINDIR", "TEMP", "TMP"):
                if value := os.environ.get(name):
                    env[name] = value

            completed = subprocess.run(
                [
                    sys.executable,
                    "-B",
                    "-c",
                    CHILD_PRELUDE + "\n" + textwrap.dedent(body),
                    str(temp_root),
                ],
                cwd=ROOT,
                env=env,
                capture_output=True,
                text=True,
                timeout=10,
                check=False,
                shell=False,
            )
            self.assertEqual(
                completed.returncode,
                0,
                msg=f"stdout:\n{completed.stdout}\nstderr:\n{completed.stderr}",
            )
            result_lines = [
                line for line in completed.stdout.splitlines() if line.startswith("RESULT=")
            ]
            self.assertEqual(len(result_lines), 1, msg=completed.stdout)
            self.assertEqual(json.loads(result_lines[0][7:]), {"ok": True})

    def test_cached_image_bytes_skip_network(self):
        self._run_case(
            r'''
            cache = temp_root / "cache" / "image.jpg"
            cache.parent.mkdir()
            cache.write_bytes(b"cached-bytes")
            service = export_service.ExportService(str(temp_root / "images"))
            with patch.object(export_service.urllib.request, "urlopen", side_effect=AssertionError("network")) as urlopen:
                result = service._fetch_image_bytes("https://images.test/a.jpg", str(cache))
            assert result == b"cached-bytes"
            urlopen.assert_not_called()
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )

    def test_remote_image_fetch_normalizes_request_and_populates_cache(self):
        self._run_case(
            r'''
            class Response:
                def __enter__(self): return self
                def __exit__(self, *args): return False
                def read(self): return b"downloaded-bytes"

            cache = temp_root / "cache" / "image.jpg"
            service = export_service.ExportService(str(temp_root / "images"))
            with patch.object(export_service.urllib.request, "urlopen", return_value=Response()) as urlopen:
                result = service._fetch_image_bytes(" http://images.test/a.jpg ", str(cache))
            request = urlopen.call_args.args[0]
            assert request.full_url == "https://images.test/a.jpg"
            assert urlopen.call_args.kwargs == {"timeout": 8}
            assert "Mozilla/5.0" in request.get_header("User-agent")
            assert request.get_header("Referer") == "https://www.xiaohongshu.com/"
            assert result == b"downloaded-bytes"
            assert cache.read_bytes() == result
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )

    def test_fetch_failure_returns_none_without_cache(self):
        self._run_case(
            r'''
            cache = temp_root / "cache" / "image.jpg"
            service = export_service.ExportService(str(temp_root / "images"))
            with patch.object(export_service.urllib.request, "urlopen", side_effect=OSError("offline")):
                result = service._fetch_image_bytes("https://images.test/a.jpg", str(cache))
            assert result is None
            assert not cache.exists()
            assert not cache.parent.exists()
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )

    def test_note_row_limits_images_to_ten_and_preserves_fallbacks(self):
        self._run_case(
            r'''
            from PIL import Image
            from openpyxl import Workbook

            png = io.BytesIO()
            Image.new("RGB", (1, 1), color="red").save(png, format="PNG")
            urls = [f"https://images.test/{index}.png" for index in range(12)]
            workbook = Workbook()
            sheet = workbook.active
            service = export_service.ExportService(str(temp_root / "images"))
            fetch = AsyncMock(side_effect=[png.getvalue()] + [None] * 9)
            service._fetch_image_bytes_threadsafe = fetch

            asyncio.run(service._write_note_row(
                sheet,
                2,
                "note-id",
                {"d_level": "D1", "quality": "A"},
                {"image_list": ",".join(urls)},
            ))

            assert fetch.await_count == 10
            calls = fetch.await_args_list
            assert [call.args[0] for call in calls] == urls[:10]
            expected = [str(temp_root / "images" / "note-id" / f"{index}.jpg") for index in range(10)]
            assert [call.args[1] for call in calls] == expected
            assert len(sheet._images) == 1
            assert [sheet.cell(2, column).value for column in range(12, 21)] == urls[1:10]
            assert sheet.cell(2, 21).value is None
            assert sheet.row_dimensions[2].height == 92
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )

    def test_note_row_sanitizes_excel_text_and_image_fallback(self):
        self._run_case(
            r'''
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            service = export_service.ExportService(str(temp_root / "images"))
            service._fetch_image_bytes_threadsafe = AsyncMock(return_value=None)
            image_url = "=HYPERLINK(\"https://example.test\")\x00"
            asyncio.run(service._write_note_row(
                sheet,
                2,
                "note-id",
                {"d_level": "D1", "quality": "A"},
                {
                    "title": "Title\x00with control",
                    "desc": "=SUM(1,2)",
                    "nickname": "+author",
                    "note_url": "https://example.test/note\x01",
                    "image_list": json.dumps([image_url]),
                },
            ))
            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)
            saved = export_service.openpyxl.load_workbook(output).active
            assert saved.cell(2, 3).value == "Titlewith control"
            assert saved.cell(2, 4).value == "'=SUM(1,2)"
            assert saved.cell(2, 5).value == "'+author"
            assert saved.cell(2, 10).value == "https://example.test/note"
            assert saved.cell(2, 11).value == "'=HYPERLINK(\"https://example.test\")"
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )

    def test_load_notes_uses_exact_parameterized_query(self):
        self._run_case(
            r'''
            calls = []
            class Cursor:
                async def fetchall(self):
                    return [{"note_id": "n2", "title": "two"}, {"note_id": "n1", "title": "one"}]
            class Connection:
                row_factory = None
                async def execute(self, sql, parameters):
                    calls.append((" ".join(sql.split()), parameters))
                    return Cursor()
            connection = Connection()
            class Context:
                async def __aenter__(self): return connection
                async def __aexit__(self, *args): return False

            export_service.aiosqlite.connect = lambda path: Context()
            export_service.SQLITE_DB_PATH = str(temp_root / "content.db")
            service = export_service.ExportService(str(temp_root / "images"))
            with patch.object(export_service.os.path, "exists", return_value=True):
                result = asyncio.run(service._load_notes(["n1", "n2", "n3"]))
            assert calls == [("SELECT * FROM xhs_note WHERE note_id IN (?,?,?)", ["n1", "n2", "n3"])]
            assert connection.row_factory is export_service.aiosqlite.Row
            assert result == {"n2": {"note_id": "n2", "title": "two"}, "n1": {"note_id": "n1", "title": "one"}}
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )

    def test_workbook_orders_tags_and_returns_rewound_xlsx(self):
        self._run_case(
            r'''
            tags = {
                "n3": {"d_level": "D2", "quality": "A"},
                "n1": {"d_level": "D1", "quality": "B"},
                "n2": {"d_level": "D1", "quality": "A"},
            }
            service = export_service.ExportService(str(temp_root / "images"))
            service._load_notes = AsyncMock(return_value={})
            service._write_note_row = AsyncMock()
            with patch.object(export_service.sdb, "get_all_note_tags", AsyncMock(return_value=tags)):
                result = asyncio.run(service.build_tagged_notes_workbook())

            assert result.tell() == 0
            assert result.read(2) == b"PK"
            assert service._load_notes.await_args.args[0] == ["n3", "n1", "n2"]
            ordered_ids = [call.args[2] for call in service._write_note_row.await_args_list]
            assert ordered_ids == ["n2", "n1", "n3"]

            result.seek(0)
            workbook = export_service.openpyxl.load_workbook(result)
            sheet = workbook["notes"]
            headers = [sheet.cell(1, column).value for column in range(1, 21)]
            assert headers == [
                "D Level", "Quality", "Title", "Description", "Author",
                "Likes", "Comments", "Collects", "Publish Date", "Note URL",
                "Image 1", "Image 2", "Image 3", "Image 4", "Image 5",
                "Image 6", "Image 7", "Image 8", "Image 9", "Image 10",
            ]
            print('RESULT=' + json.dumps({"ok": True}))
            '''
        )


if __name__ == "__main__":
    unittest.main()
