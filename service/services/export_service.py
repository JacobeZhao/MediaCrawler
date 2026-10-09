import asyncio
import io
import os
import re
import time
import urllib.request
from typing import Optional

import aiosqlite
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

from config.db_config import SQLITE_DB_PATH

from .. import service_db as sdb
from .note_service import parse_image_list


class ExportService:
    def __init__(self, image_dir: str):
        self._image_dir = image_dir

    @staticmethod
    def _safe_excel_value(value):
        if value is None:
            return ""
        if isinstance(value, str):
            value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32767]
            if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
                return "'" + value[:32766]
        return value

    async def build_selected_tasks_workbook(self, task_ids: list[int]) -> io.BytesIO:
        selected = sdb.validate_task_ids(task_ids)
        await sdb.get_selected_tasks(selected)
        if not os.path.exists(SQLITE_DB_PATH):
            raise FileNotFoundError("Selected tasks have no attributed content.")

        placeholders = ",".join("?" for _ in selected)
        async with aiosqlite.connect(SQLITE_DB_PATH) as db:
            db.row_factory = aiosqlite.Row
            cursor = await db.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='xhs_content_source'"
            )
            if await cursor.fetchone() is None:
                raise FileNotFoundError("Selected tasks have no attributed content.")
            cursor = await db.execute(
                "SELECT n.note_id, n.title, n.desc, n.nickname, n.user_id, n.time, "
                "n.liked_count, n.comment_count, n.collected_count, n.note_url, "
                "n.image_list, GROUP_CONCAT(DISTINCT s.task_id) AS task_ids "
                "FROM xhs_content_source s JOIN xhs_note n ON n.note_id=s.entity_id "
                f"WHERE s.entity_type='note' AND s.task_id IN ({placeholders}) "
                "GROUP BY n.note_id ORDER BY n.note_id",
                selected,
            )
            notes = [dict(row) for row in await cursor.fetchall()]
            cursor = await db.execute(
                "SELECT c.comment_id, c.note_id, c.content, c.nickname, c.user_id, "
                "c.create_time, c.like_count, c.parent_comment_id, c.pictures, "
                "GROUP_CONCAT(DISTINCT s.task_id) AS task_ids "
                "FROM xhs_content_source s JOIN xhs_note_comment c ON c.comment_id=s.entity_id "
                f"WHERE s.entity_type='comment' AND s.task_id IN ({placeholders}) "
                "GROUP BY c.comment_id ORDER BY c.comment_id",
                selected,
            )
            comments = [dict(row) for row in await cursor.fetchall()]

        if not notes and not comments:
            raise FileNotFoundError("Selected tasks have no attributed content.")

        wb = openpyxl.Workbook()
        notes_ws = wb.active
        notes_ws.title = "notes"
        note_fields = (
            "task_ids", "note_id", "title", "desc", "nickname", "user_id",
            "time", "liked_count", "comment_count", "collected_count",
            "note_url", "image_list",
        )
        comments_ws = wb.create_sheet("comments")
        comment_fields = (
            "task_ids", "comment_id", "note_id", "content", "nickname",
            "user_id", "create_time", "like_count", "parent_comment_id", "pictures",
        )
        for ws, fields, rows in (
            (notes_ws, note_fields, notes),
            (comments_ws, comment_fields, comments),
        ):
            ws.append(fields)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(fields))}{len(rows) + 1}"
            for row in rows:
                ws.append([self._safe_excel_value(row.get(field)) for field in fields])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    async def build_tagged_notes_workbook(self) -> io.BytesIO:
        tags = await sdb.get_all_note_tags()
        if not tags:
            raise FileNotFoundError("No tagged notes are available for export.")

        notes_data = await self._load_notes(list(tags.keys()))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "notes"

        headers = [
            "D Level", "Quality", "Title", "Description", "Author",
            "Likes", "Comments", "Collects", "Publish Date", "Note URL",
            "Image 1", "Image 2", "Image 3", "Image 4", "Image 5",
            "Image 6", "Image 7", "Image 8", "Image 9", "Image 10",
        ]
        header_fill = PatternFill("solid", fgColor="F43F5E")
        header_font = Font(bold=True, color="FFFFFF")
        ws.row_dimensions[1].height = 20
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sorted_ids = sorted(tags.keys(), key=lambda nid: (tags[nid]["d_level"], tags[nid]["quality"]))
        for row_idx, note_id in enumerate(sorted_ids, 2):
            tag = tags[note_id]
            note = notes_data.get(note_id, {})
            await self._write_note_row(ws, row_idx, note_id, tag, note)

        col_widths = [8, 12, 30, 60, 15, 8, 8, 8, 12, 40] + [18] * 10
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    async def _load_notes(self, note_ids: list[str]) -> dict:
        if not os.path.exists(SQLITE_DB_PATH):
            return {}
        notes_data = {}
        async with aiosqlite.connect(SQLITE_DB_PATH) as db:
            db.row_factory = aiosqlite.Row
            placeholders = ",".join("?" * len(note_ids))
            cursor = await db.execute(
                f"SELECT * FROM xhs_note WHERE note_id IN ({placeholders})", note_ids
            )
            for row in await cursor.fetchall():
                notes_data[row["note_id"]] = dict(row)
        return notes_data

    async def _write_note_row(self, ws, row_idx: int, note_id: str, tag: dict, note: dict):
        imgs = parse_image_list(note.get("image_list"))
        pub_time = ""
        ts = note.get("time")
        if ts:
            try:
                pub_time = time.strftime("%Y-%m-%d", time.localtime(int(ts) / 1000))
            except Exception:
                pub_time = str(ts)

        values = [
            tag["d_level"],
            tag["quality"],
            note.get("title", ""),
            note.get("desc", ""),
            note.get("nickname", ""),
            note.get("liked_count", ""),
            note.get("comment_count", ""),
            note.get("collected_count", ""),
            pub_time,
            note.get("note_url", f"https://www.xiaohongshu.com/explore/{note_id}"),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=self._safe_excel_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 4))

        ws.row_dimensions[row_idx].height = 92
        for img_idx, url in enumerate(imgs[:10]):
            col_idx = 11 + img_idx
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            local_path = os.path.join(self._image_dir, note_id, f"{img_idx}.jpg")
            img_bytes = await self._fetch_image_bytes_threadsafe(url, local_path)
            if not img_bytes:
                ws.cell(row=row_idx, column=col_idx, value=self._safe_excel_value(url))
                continue
            jpeg_bytes = self._to_jpeg(img_bytes) or img_bytes
            try:
                xl_img = XLImage(io.BytesIO(jpeg_bytes))
                xl_img.width = 120
                xl_img.height = 120
                ws.add_image(xl_img, f"{col_letter}{row_idx}")
            except Exception:
                ws.cell(row=row_idx, column=col_idx, value=self._safe_excel_value(url))

    async def _fetch_image_bytes_threadsafe(self, url: str, local_path: str) -> Optional[bytes]:
        return await asyncio.to_thread(self._fetch_image_bytes, url, local_path)

    def _fetch_image_bytes(self, url: str, local_path: str) -> Optional[bytes]:
        url = url.strip().strip('"').strip("'")
        if url.startswith("http://"):
            url = "https://" + url[7:]
        if os.path.exists(local_path):
            with open(local_path, "rb") as f:
                return f.read()
        try:
            req = urllib.request.Request(
                url,
                headers={
                    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
                    "Referer": "https://www.xiaohongshu.com/",
                },
            )
            with urllib.request.urlopen(req, timeout=8) as resp:
                raw = resp.read()
            if raw:
                os.makedirs(os.path.dirname(local_path), exist_ok=True)
                with open(local_path, "wb") as f:
                    f.write(raw)
            return raw
        except Exception:
            return None

    def _to_jpeg(self, raw: bytes) -> Optional[bytes]:
        try:
            from PIL import Image as PILImage

            img = PILImage.open(io.BytesIO(raw))
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="JPEG", quality=85)
            return buf.getvalue()
        except Exception:
            return None
